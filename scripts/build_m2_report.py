#!/usr/bin/env python3
"""Build M2 diagrams and a Word report from the reviewed Markdown source.

Requirements: python-docx, Pillow, a Chinese font.
Run from any directory: python scripts/build_m2_report.py
Render separately with render_docx.py, LibreOffice and Poppler.
"""
from pathlib import Path
import html
import math
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs/requirements/assets/m2"
FONT = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
BOLD = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")
INK = "#172C45"
BLUE = "#245D8D"
TEAL = "#16736C"
GRAY = "#5F6E7C"
SCALE = 3

class Scene:
    def __init__(self, width, height):
        self.width, self.height = width, height
        self.img = Image.new("RGB", (width*SCALE, height*SCALE), "white")
        self.draw = ImageDraw.Draw(self.img)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                    '<rect width="100%" height="100%" fill="white"/>']
    def text(self, x, y, value, size=24, bold=False, color=INK, center=True):
        font = ImageFont.truetype(str(BOLD if bold else FONT), round(size*SCALE), index=2)
        self.draw.text((x*SCALE,y*SCALE),value,font=font,fill=color,anchor="mm" if center else "lm")
        anchor = "middle" if center else "start"
        weight = "700" if bold else "400"
        self.svg.append(f'<text x="{x}" y="{y}" dominant-baseline="central" text-anchor="{anchor}" font-family="Noto Sans CJK SC,Microsoft YaHei,sans-serif" font-size="{size}" font-weight="{weight}" fill="{color}">{html.escape(value)}</text>')
    def box(self, x, y, w, h, title, subtitle=None, color=BLUE):
        fill = "#F1F6FB" if color == BLUE else "#EEF7F5"
        rect = (x*SCALE,y*SCALE,(x+w)*SCALE,(y+h)*SCALE)
        self.draw.rounded_rectangle(rect,radius=12*SCALE,fill=fill,outline=color,width=2*SCALE)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="12" fill="{fill}" stroke="{color}" stroke-width="2"/>')
        lines = [] if subtitle is None else subtitle.split("\n")
        title_y = y+h/2 if not lines else y+(26 if h<=110 else 30)
        self.text(x+w/2,title_y,title,27,True,color)
        for i,line in enumerate(lines):
            sub_y = y+58+i*26 if h<=110 else y+70+i*30
            self.text(x+w/2,sub_y,line,20 if len(lines)>1 and h<=110 else 22,False,INK)
    def line(self, pts, color=GRAY, dashed=False, arrow=True):
        pixels=[(round(x*SCALE),round(y*SCALE)) for x,y in pts]
        self.draw.line(pixels,fill=color,width=3*SCALE,joint="curve")
        coords = " ".join(f"{x},{y}" for x,y in pts)
        dash = ' stroke-dasharray="8 6"' if dashed else ""
        self.svg.append(f'<polyline points="{coords}" fill="none" stroke="{color}" stroke-width="3" stroke-linejoin="round"{dash}/>')
        if arrow:
            x,y=pts[-1]; px,py=pts[-2]
            angle=math.atan2(y-py,x-px)
            tri=[(x,y),(x-13*math.cos(angle)-6*math.sin(angle),y-13*math.sin(angle)+6*math.cos(angle)),(x-13*math.cos(angle)+6*math.sin(angle),y-13*math.sin(angle)-6*math.cos(angle))]
            self.draw.polygon([(round(a*SCALE),round(b*SCALE)) for a,b in tri],fill=color)
            self.svg.append('<polygon points="'+" ".join(f"{a:.1f},{b:.1f}" for a,b in tri)+f'" fill="{color}"/>')
    def save(self, stem):
        ASSETS.mkdir(parents=True,exist_ok=True)
        self.img.save(ASSETS/f"{stem}.png",dpi=(300,300))
        (ASSETS/f"{stem}.svg").write_text("\n".join(self.svg+["</svg>"]),encoding="utf-8")

def diagrams():
    s=Scene(1160,750)
    s.text(265,25,"记录与保存",25,True,TEAL)
    s.text(875,25,"提问与回答",25,True,BLUE)
    s.box(40,60,450,70,"长期对话 Dialogue",color=TEAL)
    s.box(40,170,450,110,"Memory Construction","事实/事件抽取 · 人物/时间整理\n来源与字段检查",TEAL)
    s.box(40,335,450,150,"Memory Repository","Structured Memory / History\n状态关系 · Current State\n来源证据",TEAL)
    s.box(40,585,450,85,"原文核查与 JSON 导出",color=TEAL)
    s.line([(265,130),(265,170)],TEAL)
    s.line([(265,280),(265,335)],TEAL)
    s.line([(265,485),(265,585)],TEAL)
    s.box(650,60,450,70,"用户问题 Question")
    s.box(650,170,450,110,"Hybrid Retrieval","BM25 + Vector + Entity/Time\n按问题对象与时间寻找候选")
    s.box(650,325,450,70,"相关 Memory")
    s.box(650,435,450,110,"State Resolution","历史 + 关系 + 问题时间\n冲突无法消解时保留不确定性")
    s.box(650,585,450,85,"有效 Memory → Answer","Web 展示答案、证据及状态过程")
    for a,b in [(130,170),(280,325),(395,435),(545,585)]:
        s.line([(875,a),(875,b)],BLUE)
    s.line([(490,355),(550,355),(550,225),(650,225)],TEAL)
    s.text(554,307,"历史",19,False,TEAL)
    s.line([(490,465),(570,465),(570,490),(650,490)],TEAL)
    s.text(566,427,"关系与状态",18,False,TEAL)
    s.text(580,718,"保留旧 Memory；有效证据由问题与时间条件决定",23,True,INK)
    s.save("system-architecture")

    s=Scene(1160,860)
    s.text(285,28,"Memory Recording",29,True,TEAL)
    s.text(875,28,"Memory Retrieval",29,True,BLUE)
    recording=[
        ("Dialogue","课程对话与原始来源"),
        ("Memory Extraction","大模型提取重要事实和事件"),
        ("人物、时间与来源整理","模型辅助理解 · 规则检查"),
        ("Structured Memory","保存内容、人物、时间与来源"),
        ("关系判断与状态维护","RETAIN / UPDATE / SUPERSEDE\n无法消解的 CONFLICT"),
        ("History + Current State","旧 Memory 保留，支持重复查询"),
    ]
    retrieval=[
        ("Question","识别问题对象与时间条件"),
        ("Hybrid Retrieval","BM25 + Vector + Entity/Time"),
        ("Relevant Memories","合并候选、去重、保留来源"),
        ("State Resolution","历史关系 + Current State\n按问题时间判断适用性"),
        ("Effective Memory","选择有效证据或说明不确定性"),
        ("Answer + Evidence","展示原文及状态解析过程"),
    ]
    for col,steps,color in [(55,recording,TEAL),(645,retrieval,BLUE)]:
        for i,(title,sub) in enumerate(steps):
            y=70+125*i
            s.box(col,y,460,100,title,sub,color)
            if i<5:
                s.line([(col+230,y+100),(col+230,y+125)],color)
    s.text(580,826,"当前问题、历史问题与未消解冲突采用不同证据处理",23,True)
    s.save("recording-retrieval-flow")

def set_font(run,size=11,bold=None):
    run.font.name="Noto Sans CJK SC"
    run.font.size=Pt(size)
    run.font.color.rgb=RGBColor(0,0,0)
    if bold is not None:
        run.bold=bold
    fonts=run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("ascii","hAnsi","eastAsia","cs"):
        fonts.set(qn("w:"+k),"Noto Sans CJK SC")

def text_runs(p,text,size=11):
    # Render source links as clickable Word hyperlinks, retaining readable labels.
    text=text.replace("`","")
    start=0
    for match in re.finditer(r"\[([^\]]+)\]\((https?://[^)]+)\)",text):
        set_font(p.add_run(text[start:match.start()]),size)
        link=OxmlElement("w:hyperlink")
        link.set(qn("r:id"),p.part.relate_to(match.group(2),RT.HYPERLINK,is_external=True))
        run=p.add_run(match.group(1));set_font(run,size)
        link.append(run._element);p._p.append(link)
        start=match.end()
    set_font(p.add_run(text[start:]),size)
    return p

def table(doc,rows):
    widths= [4.2,12.4] if len(rows[0])==2 else [3.6,7.6,5.4]
    t=doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    borders=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,w in enumerate(widths):
        t.columns[i].width=Cm(w)
    for rindex,data in enumerate(rows):
        row=t.rows[0] if rindex==0 else t.add_row()
        pr=row._tr.get_or_add_trPr()
        cant=OxmlElement("w:cantSplit");pr.append(cant)
        if rindex==0:
            repeat=OxmlElement("w:tblHeader");pr.append(repeat)
        for col,value in enumerate(data):
            cell=row.cells[col]
            cell.width=Cm(widths[col])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement("w:shd")
            shade.set(qn("w:fill"),"E7EDF4" if rindex==0 else ("F6F8FA" if rindex%2==0 else "FFFFFF"))
            tcpr.append(shade)
            margins=OxmlElement("w:tcMar")
            for edge,val in (("top","65"),("bottom","65"),("left","95"),("right","95")):
                e=OxmlElement("w:"+edge);e.set(qn("w:w"),val);e.set(qn("w:type"),"dxa");margins.append(e)
            tcpr.append(margins)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.08
            r=p.add_run(value)
            set_font(r,9.5 if len(rows[0])==3 else 10, rindex==0)
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(1)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(2)
    spacer.add_run().font.size=Pt(2)
    return t

def build():
    diagrams()
    source=ROOT/"docs/requirements/vision-and-scope.md"
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.65)
    sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    sec.header_distance=Cm(0.65);sec.footer_distance=Cm(0.7)
    sec.different_first_page_header_footer=False
    doc.settings.odd_and_even_pages_header_footer=False
    for element in doc.styles.element.xpath('.//w:pBdr'):
        element.getparent().remove(element)
    for style in ["Normal","Title","Subtitle","Heading 1","Heading 2","Caption"]:
        st=doc.styles[style]
        st.font.name="Noto Sans CJK SC";st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"Noto Sans CJK SC")
        st.paragraph_format.space_after=Pt(5)
    doc.styles["Normal"].font.size=Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing=1.15
    doc.styles["Title"].font.size=Pt(21)
    doc.styles["Title"].font.bold=True
    doc.styles["Title"].paragraph_format.space_after=Pt(8)
    doc.styles["Heading 1"].font.size=Pt(13)
    doc.styles["Heading 1"].font.bold=True
    doc.styles["Heading 1"].paragraph_format.space_before=Pt(9)
    doc.styles["Heading 1"].paragraph_format.space_after=Pt(5)
    doc.styles["Heading 1"].paragraph_format.keep_with_next=True
    doc.styles["Caption"].font.size=Pt(9)
    doc.styles["Caption"].font.bold=False
    doc.styles["Caption"].paragraph_format.space_after=Pt(6)
    header=sec.header.paragraphs[0]
    text_runs(header,"在阴雨天游泳   |   Team 3   |   M2 项目提案",8)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    text_runs(footer,"Vision and Scope    ·    ",8)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    lines=source.read_text(encoding="utf-8").splitlines()
    i=0
    pending_pagebreak=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:
            i+=1;continue
        if line=="<!-- pagebreak -->":
            pending_pagebreak=True
        elif line.startswith("# "):
            p=doc.add_paragraph(style="Title");text_runs(p,line[2:],21);p.runs[0].bold=True
        elif line.startswith("## "):
            p=doc.add_paragraph(style="Heading 1");text_runs(p,line[3:],13);p.runs[0].bold=True
            if pending_pagebreak:
                p.paragraph_format.page_break_before=True
                pending_pagebreak=False
        elif line.startswith("!["):
            match=re.match(r"!\[(.*?)\]\((.*?)\)",line)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            r=p.add_run();r.add_picture(str(source.parent/match.group(2)),width=Cm(16.6))
            for pic in r._element.xpath(".//wp:docPr"):
                pic.set("descr",match.group(1))
        elif line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                vals=[s.strip() for s in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",s.replace(" ","")) for s in vals):
                    rows.append(vals)
                i+=1
            table(doc,rows)
            continue
        elif line.startswith("图 "):
            p=doc.add_paragraph(style="Caption");text_runs(p,line,9)
        else:
            p=doc.add_paragraph()
            p.paragraph_format.widow_control=True
            text_runs(p,line,9 if re.match(r"\[\d+\]",line) else (10 if line.startswith(("Milestone","队名：","成员：")) else 11))
        i+=1
    doc.core_properties.title="动态对话记忆管理项目提案"
    doc.core_properties.subject="M2 Vision and Scope and Feature Roadmap"
    doc.core_properties.author="在阴雨天游泳项目组"
    doc.core_properties.keywords="M2, LoCoMo, Dynamic Memory"
    doc.core_properties.comments="本地制作稿，正式交付前按团队流程复核。"
    output=source.with_suffix(".docx")
    doc.save(output)
    print(f"Created {output}")
    print("Created 2 diagrams in PNG and SVG formats")

if __name__=="__main__":
    build()
