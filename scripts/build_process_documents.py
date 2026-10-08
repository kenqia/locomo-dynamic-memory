#!/usr/bin/env python3
"""Render M1 meeting, snapshot, and member records as editable Word files."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from build_m2_report import ROOT, text_runs

SOURCES = [
    'process/meetings/first-meeting.md',
    'process/snapshots/2026-10-06-m1-m2.md',
    'process/snapshots/2026-10-06-review-sync.md',
    'process/reviews/m1-member-review.md',
]

def build(path):
    source = ROOT / path
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    for name in ['Normal', 'Title', 'Heading 1']:
        style = doc.styles[name]
        style.font.name = 'Noto Sans CJK SC'
        style.font.size = Pt(10 if name == 'Normal' else 16 if name == 'Title' else 12)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        style.paragraph_format.space_after = Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                values = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in values):
                    rows.append(values)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            for index, values in enumerate(rows):
                row = table.add_row()
                pr = row._tr.get_or_add_trPr()
                pr.append(OxmlElement('w:cantSplit'))
                if index == 0:
                    pr.append(OxmlElement('w:tblHeader'))
                for cell, value in zip(row.cells, values):
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    text_runs(p, value, 9)
                    if index == 0:
                        for run in p.runs:
                            run.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            text_runs(p, line[2:], 16)
        elif line.startswith('## '):
            p = doc.add_paragraph(style='Heading 1')
            text_runs(p, line[3:], 12)
            if line.startswith('## 会议后进展'):
                p.paragraph_format.page_break_before = True
        else:
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            text_runs(p, line, 10)
        i += 1
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text_runs(footer, '在阴雨天游泳 · 过程材料 · ', 8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = lines[0].lstrip('# ')
    doc.core_properties.author = '在阴雨天游泳项目组'
    doc.save(source.with_suffix('.docx'))
    print(source.with_suffix('.docx').relative_to(ROOT))

if __name__ == '__main__':
    for source in SOURCES:
        build(source)
