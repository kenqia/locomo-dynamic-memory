"""Build the M4 Markdown source into a five-page PDF and editable Word file.

Document tools only: Pillow, python-docx, reportlab and PyMuPDF.
Install them in a separate environment; they are not application dependencies.
Chinese fonts default to Windows SimSun/SimHei, or pass --font and --bold-font.
"""
from __future__ import annotations

import argparse
from hashlib import sha256
from html import escape
import json
import os
from pathlib import Path
import re
from urllib.parse import unquote, urlsplit

from PIL import Image as PillowImage, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import pymupdf
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Flowable, Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/design/m4-midterm-report.md"
ASSETS = SOURCE.parent / "assets/m4"
OUTPUT = ROOT / "deliverables/m4"
STEM = "PROJECT-MILESTONE4-3-在阴雨天游泳-报告"
INK = "#18334D"
BLUE = "#245D8D"


def architecture(font_path: Path) -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    picture = PillowImage.new("RGB", (1500, 480), "white")
    draw = ImageDraw.Draw(picture)
    font = ImageFont.truetype(str(font_path), 27)
    boxes = [
        (20, 185, 210, 285, ["Streamlit", "交互与展示"]),
        (265, 185, 465, 285, ["FastAPI", "公开业务接口"]),
        (535, 20, 965, 125, ["构建 / 关系 / 状态投影"]),
        (535, 185, 965, 290, ["检索 / 扩展 / 解析 / 回答"]),
        (535, 350, 965, 455, ["版本 / 导出 / 评测"]),
        (1050, 20, 1480, 125, ["DeepSeek 模型适配器"]),
        (1050, 185, 1480, 290, ["SQLite 来源与版本产物"]),
        (1050, 350, 1480, 455, ["BM25 / MiniLM 索引"]),
    ]
    lines = [((210, 235), (265, 235)), ((465, 235), (505, 235)),
             ((505, 72), (505, 402)), ((505, 72), (535, 72)),
             ((505, 235), (535, 235)), ((505, 402), (535, 402)),
             ((965, 72), (1005, 72)), ((965, 235), (1005, 235)),
             ((965, 402), (1005, 402)), ((1005, 72), (1005, 402)),
             ((1005, 72), (1050, 72)), ((1005, 235), (1050, 235)),
             ((1005, 402), (1050, 402))]
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1500" height="480" viewBox="0 0 1500 480">',
           '<rect width="1500" height="480" fill="white"/>']
    for start, end in lines:
        draw.line([start, end], fill=BLUE, width=4)
        svg.append(f'<line x1="{start[0]}" y1="{start[1]}" x2="{end[0]}" y2="{end[1]}" stroke="{BLUE}" stroke-width="4"/>')
    for left, top, right, bottom, labels in boxes:
        draw.rounded_rectangle((left, top, right, bottom), radius=12, fill="#EDF4FA", outline=BLUE, width=3)
        svg.append(f'<rect x="{left}" y="{top}" width="{right-left}" height="{bottom-top}" rx="12" fill="#EDF4FA" stroke="{BLUE}" stroke-width="3"/>')
        for index, label in enumerate(labels):
            x = (left + right) / 2
            y = (top + bottom) / 2 + (index - (len(labels)-1)/2) * 36
            draw.text((x, y), label, font=font, fill=INK, anchor="mm")
            svg.append(f'<text x="{x}" y="{y}" text-anchor="middle" dominant-baseline="central" font-family="SimHei,Microsoft YaHei,sans-serif" font-size="27" fill="{INK}">{escape(label)}</text>')
    picture.save(ASSETS / "system-architecture.png")
    (ASSETS / "system-architecture.svg").write_text("\n".join(svg + ["</svg>"]), encoding="utf-8")


def blocks(text: str):
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line == "<!-- pagebreak -->":
            yield "break", None
        elif line.startswith("|"):
            rows = []
            while True:
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                if index == len(lines) or not lines[index].strip().startswith("|"):
                    break
                line = lines[index].strip()
                index += 1
            yield "table", rows
        elif match := re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line):
            target = urlsplit(match.group(2))
            view = None
            if target.fragment:
                assert target.fragment.startswith("view="), "Unknown image viewport"
                view = tuple(int(part) for part in target.fragment[5:].split(","))
                assert len(view) == 4
            yield "image", (match.group(1), SOURCE.parent / unquote(target.path), view)
        elif line.startswith("# "):
            yield "title", line[2:]
        elif line.startswith("## "):
            yield "heading", line[3:]
        elif line.startswith("### "):
            yield "subheading", line[4:]
        else:
            yield "text", line


def plain(text: str) -> str:
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text).replace("`", "")


def workflow(font_path: Path) -> None:
    """Draw the two documented workflows; the architecture remains editable code."""
    picture = PillowImage.new("RGB", (1500, 350), "white")
    draw = ImageDraw.Draw(picture)
    font = ImageFont.truetype(str(font_path), 25)
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1500" height="350" viewBox="0 0 1500 350">',
           '<rect width="1500" height="350" fill="white"/>']
    flows = [
        (20, "构建", [["导入原文", "来源检查"], ["分片抽取", "原子校验"], ["索引 / 关系", "状态投影"],
                       ["完整性检查", "发布 Build"], ["失败保留旧版", "显式恢复"]]),
        (185, "查询", [["固定 Build", "问题分析"], ["BM25 / Dense", "RRF 混合召回"], ["扩展相关历史", "选择有效证据"],
                        ["有界原文补查", "仅在需要时"], ["答案 / 原文", "保存 Trace"]]),
    ]
    for top, title, boxes in flows:
        draw.text((15, top + 48), title, font=font, fill=INK, anchor="lm")
        svg.append(f'<text x="15" y="{top+48}" font-family="SimHei,sans-serif" font-size="25" fill="{INK}">{title}</text>')
        for index, labels in enumerate(boxes):
            left = 100 + index * 280
            right = left + 250
            bottom = top + 110
            draw.rounded_rectangle((left, top, right, bottom), radius=10, fill="#EDF4FA", outline=BLUE, width=3)
            svg.append(f'<rect x="{left}" y="{top}" width="250" height="110" rx="10" fill="#EDF4FA" stroke="{BLUE}" stroke-width="3"/>')
            for line, label in enumerate(labels):
                x, y = (left + right) / 2, top + 37 + 37 * line
                draw.text((x, y), label, font=font, fill=INK, anchor="mm")
                svg.append(f'<text x="{x}" y="{y}" text-anchor="middle" dominant-baseline="central" font-family="SimHei,sans-serif" font-size="25" fill="{INK}">{escape(label)}</text>')
            if index < 4:
                # Build recovery is a side branch of the build, not a step after publication.
                if title == "构建" and index == 3:
                    continue
                x, y = right + 25, top + 55
                draw.line([(right + 3, y), (x, y)], fill=BLUE, width=3)
                draw.polygon([(x, y), (x-9, y-6), (x-9, y+6)], fill=BLUE)
                svg.append(f'<path d="M {right+3} {y} H {x} M {x-9} {y-6} L {x} {y} L {x-9} {y+6}" fill="none" stroke="{BLUE}" stroke-width="3"/>')
    note = "构建写入版本与全局状态；查询只读完整版本，仅保存本次 Trace。"
    draw.text((750, 330), note, font=font, fill=INK, anchor="mm")
    svg.append(f'<text x="750" y="330" text-anchor="middle" font-family="SimHei,sans-serif" font-size="25" fill="{INK}">{note}</text>')
    picture.save(ASSETS / "build-query-flow.png")
    (ASSETS / "build-query-flow.svg").write_text("\n".join(svg + ["</svg>"]), encoding="utf-8")


class ScreenshotViewport(Flowable):
    """Clip the image in the document layout while preserving the source bitmap."""
    def __init__(self, path: Path, size: tuple[int, int], view: tuple[int, ...], width: float):
        super().__init__()
        self.path, self.size, self.view = path, size, view
        left, top, right, bottom = view
        assert 0 <= left < right <= size[0] and 0 <= top < bottom <= size[1]
        self.width = width
        self.height = width * (bottom-top) / (right-left)

    def draw(self):
        left, top, right, bottom = self.view
        scale = self.width / (right-left)
        self.canv.saveState()
        clipping = self.canv.beginPath()
        clipping.rect(0, 0, self.width, self.height)
        self.canv.clipPath(clipping, stroke=0)
        self.canv.drawImage(str(self.path), -left*scale, -(self.size[1]-bottom)*scale,
                            width=self.size[0]*scale, height=self.size[1]*scale)
        self.canv.restoreState()


def inline(text: str, baseline: str) -> str:
    def link(match):
        target = (SOURCE.parent / unquote(match.group(2))).resolve()
        relative = target.relative_to(ROOT).as_posix()
        if relative.startswith("process/m4/"):
            url = Path(os.path.relpath(target, OUTPUT)).as_posix()
        else:
            url = f"https://github.com/kenqia/locomo-dynamic-memory/blob/{baseline}/{relative}"
        return f'<link href="{escape(url, quote=True)}" color="{BLUE}">{escape(match.group(1))}</link>'
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, escape(text).replace("`", ""))


def build(font: Path, bold_font: Path, reference_baseline: str = "cf422fe"):
    text = SOURCE.read_text(encoding="utf-8")
    baseline = re.search(r"main@([a-f0-9]+)", text).group(1)
    reference_match = re.search(r"文档引用基准：main@([a-f0-9]+)", text)
    if reference_match:
        reference_baseline = reference_match.group(1)
    revision = re.search(r"版本：(v[0-9.]+)", text).group(1)
    assert text.count("<!-- pagebreak -->") == 4, "Source must have five explicit page sections"
    architecture(bold_font)
    workflow(bold_font)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(TTFont("M4Body", str(font), subfontIndex=0))
    pdfmetrics.registerFont(TTFont("M4Bold", str(bold_font), subfontIndex=0))
    styles = {
        "text": ParagraphStyle("Body", fontName="M4Body", fontSize=10.5, leading=16.5, wordWrap="CJK", spaceAfter=7),
        "title": ParagraphStyle("Title", fontName="M4Bold", fontSize=18, leading=26, spaceAfter=10, textColor=colors.HexColor(INK)),
        "heading": ParagraphStyle("Heading", fontName="M4Bold", fontSize=14, leading=20, spaceAfter=10, keepWithNext=True, textColor=colors.HexColor(INK)),
        "subheading": ParagraphStyle("Subheading", fontName="M4Bold", fontSize=11.5, leading=18, spaceBefore=5, spaceAfter=6, keepWithNext=True),
        "cell": ParagraphStyle("Cell", fontName="M4Body", fontSize=9.4, leading=14, wordWrap="CJK"),
        "caption": ParagraphStyle("Caption", fontName="M4Body", fontSize=9, leading=13, wordWrap="CJK", spaceAfter=8),
    }
    pdf = OUTPUT / f"{STEM}.pdf"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in (("Normal", 10.5), ("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "SimSun"
        style.font.size = Pt(size)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1
    doc.core_properties.title = "动态对话记忆管理：M4 中期报告"
    doc.core_properties.author = "在阴雨天游泳项目组"
    story = []
    width = A4[0] - 40 * mm
    for kind, value in blocks(text):
        if kind == "break":
            story.append(PageBreak())
            doc.add_page_break()
        elif kind == "image":
            with PillowImage.open(value[1]) as image:
                size = image.size
            view = value[2]
            if view:
                figure = ScreenshotViewport(value[1], size, view, width)
                shape = doc.add_picture(str(value[1]), width=Cm(17), height=Cm(17*(view[3]-view[1])/(view[2]-view[0])))
                crop = OxmlElement("a:srcRect")
                for name, fraction in zip(("l", "t", "r", "b"),
                        (view[0]/size[0], view[1]/size[1], 1-view[2]/size[0], 1-view[3]/size[1])):
                    crop.set(name, str(round(fraction*100000)))
                shape._inline.graphic.graphicData.pic.blipFill.insert(1, crop)
            else:
                figure = Image(str(value[1]), width=width, height=width*size[1]/size[0])
                doc.add_picture(str(value[1]), width=Cm(17))
            story.extend([figure, Spacer(1, 7)])
        elif kind == "table":
            columns = len(value[0])
            fractions = {2: [0.25, 0.75], 3: [0.19, 0.34, 0.47], 4: [0.43, 0.19, 0.16, 0.22]}[columns]
            if value[0][0] == "K":
                fractions = [0.10, 0.30, 0.30, 0.30]
            elif value[0][0] == "用户任务":
                fractions = [0.16, 0.27, 0.25, 0.32]
            elif value[0][0] == "策略":
                fractions = [0.37, 0.15, 0.26, 0.22]
            elif value[0][0] == "历史阶段":
                fractions = [0.20, 0.35, 0.45]
            elif value[0][0] == "未完成项":
                fractions = [0.25, 0.75] if columns == 2 else [0.21, 0.28, 0.51]
            rows = [[Paragraph(inline(cell, reference_baseline), styles["cell"]) for cell in row] for row in value]
            table = Table(rows, colWidths=[width * part for part in fractions], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EFF6")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D4DF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 10)])
            word_table = doc.add_table(rows=0, cols=columns)
            word_table.style = "Table Grid"
            for row in value:
                cells = word_table.add_row().cells
                for cell, content in zip(cells, row):
                    cell.text = plain(content)
        else:
            style = "caption" if kind == "text" and re.match(r"图 \d", value) else kind
            story.append(Paragraph(inline(value, reference_baseline), styles[style]))
            word_style = {"title": "Title", "heading": "Heading 1", "subheading": "Heading 2"}.get(kind)
            doc.add_paragraph(plain(value), style=word_style)
    def footer(canvas, document):
        canvas.setTitle(f"动态对话记忆管理 M4 中期报告 {revision}")
        canvas.setAuthor("在阴雨天游泳项目组")
        canvas.setFont("M4Body", 8)
        canvas.setFillColor(colors.HexColor("#53687C"))
        canvas.drawString(20 * mm, 285 * mm, f"Team 3 | M4 中期报告 | {revision} 待复核")
        canvas.drawRightString(190 * mm, 10 * mm, f"第 {document.page} 页 / 共 5 页")
    SimpleDocTemplate(str(pdf), pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
                      topMargin=19*mm, bottomMargin=17*mm).build(story, onFirstPage=footer, onLaterPages=footer)
    word = SOURCE.with_suffix(".docx")
    doc.save(word)
    with pymupdf.open(pdf) as rendered:
        page_count = len(rendered)
        if page_count != 5:
            raise RuntimeError(f"Expected five pages including references; got {page_count}. Adjust source/layout.")
        preview = ROOT / "process/m4/preview"
        preview.mkdir(parents=True, exist_ok=True)
        characters = []
        for index, page in enumerate(rendered):
            page.get_pixmap(matrix=pymupdf.Matrix(1.15, 1.15)).save(preview / f"page-{index+1}.png")
            page_text = page.get_text()
            characters.append(len(page_text))
            if not re.search(rf"{index+1} ", page_text):
                raise RuntimeError("Page number missing")
        (ROOT / "process/m4/pdf-text.txt").write_text("\n\f\n".join(p.get_text() for p in rendered), encoding="utf-8")
    manifest = {"source": SOURCE.relative_to(ROOT).as_posix(), "baseline": baseline,
                "reference_baseline": reference_baseline,
                "pdf": pdf.relative_to(ROOT).as_posix(), "pages": page_count, "page_limit": 5,
                "limit_passed": page_count <= 5, "pdf_sha256": sha256(pdf.read_bytes()).hexdigest(),
                "source_sha256": sha256(SOURCE.read_bytes()).hexdigest(),
                "characters_per_page": characters,
                "word": word.relative_to(ROOT).as_posix(), "word_pagination": "not independently rendered; use verified PDF for page limit",
                "font": font.name, "bold_font": bold_font.name}
    (ROOT / "process/m4/pdf-validation.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, default=Path("C:/Windows/Fonts/simsun.ttc"))
    parser.add_argument("--bold-font", type=Path, default=Path("C:/Windows/Fonts/simhei.ttf"))
    parser.add_argument("--reference-baseline", default="cf422fe", help="Git revision used for frozen source links")
    options = parser.parse_args()
    build(options.font, options.bold_font, options.reference_baseline)
